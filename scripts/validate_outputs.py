from pathlib import Path
import json,zipfile
from bs4 import BeautifulSoup as Soup
from docx import Document
from lxml import etree
import pymupdf

plan=json.loads(Path('plan.json').read_text())
html=Soup(Path('index.html').read_text(),'html.parser')
assert len(html.select('.sheet'))==4
blank=0
for n,sheet in enumerate(html.select('.sheet')):
 template=Soup(Path('templates/'+('first-half.html' if n%2==0 else 'second-half.html')).read_text(),'html.parser')
 for j,section in enumerate(sheet.select('section')[:3]):
  cells=section.select_one('div.grid').find_all('div',recursive=False)
  original=template.select('section')[j].select_one('div.grid').find_all('div',recursive=False)
  count=len(plan[(n//2)*6+(n%2)*3+j]['boxes'])
  for k in range(count,6):
   assert str(cells[k])==str(original[k]);blank+=1
assert blank==18
assert '۵۹–۹۰' not in Path('index.html').read_text()
physics=[b for d in plan for b in d['boxes'] if b['subject']=='فیزیک دهم']
assert len(physics)==6
assert '۲۳–۵۲' in physics[0]['topic']
assert not any('نانو' in b['topic'] for b in physics)
pdf=pymupdf.open('plan-1-aban-1405.pdf')
assert len(pdf)==4
for p in pdf:assert len(p.get_text())>100

doc=Document('plan-1-aban-1405.docx')
study=[t for t in doc.tables if t.style.name=='Planner Day']
assert len(study)==12 and len(doc.tables)==18
word_blank=0
for d,t in zip(plan,study):
 cells=[c for row in [t.rows[1],t.rows[3]] for c in row.cells]
 metrics=[c for row in [t.rows[2],t.rows[4]] for c in row.cells]
 for row in t.rows:
  assert row.height_rule==2, 'Word row height must be exact'
 assert len(cells)==6
 for i,b in enumerate(d['boxes']):
  for field in ['subject','topic','detail']:assert b[field] in cells[i].text
 for i in range(len(d['boxes']),6):
  assert '☐ درس و مبحث:' in cells[i].text and 'ساعت: ..............' in metrics[i].text
  word_blank+=1
assert word_blank==18
ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
with zipfile.ZipFile('plan-1-aban-1405.docx') as z:
 xml=etree.fromstring(z.read('word/document.xml'))
 assert len(xml.xpath('//w:br[@w:type="page"]',namespaces=ns))==3
 assert len(xml.xpath('//w:tblPr/w:bidiVisual',namespaces=ns))==18
 assert not any(x.startswith('word/media/') for x in z.namelist())
 assert len([x for x in z.namelist() if x.endswith('.odttf')])==2
 assert not xml.xpath('//w:pPr/w:jc[@w:val="right"]',namespaces=ns)
print('PASS: 4 PDF pages; new physics source; 18 original HTML blank cells; editable RTL DOCX content/ordering and 18 blank cells.')
