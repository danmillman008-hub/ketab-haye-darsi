"""Shared Stitch14 branding; only logo/social icons are images, never form content."""
from docx.shared import Mm, Pt
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml.ns import qn

INSTITUTE='موسسه مشاوره کنکور رتبه لند'
TAGLINE='آکادمی مشاوره تخصصی کنکور رتبه لند - برنامه پایش اختصاصی'


def header(w,style,title,subtitle,week,date,planner=False):
 widths=[16,103,31,52] if planner else [16,112,31,39]
 h=w.grid(2 if planner else 1,4,style,widths)
 h.rows[0].height=Mm(18)
 for c in h.rows[0].cells:
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for edge in ['top','left','right']:w.border(c,edge,val='nil')
  w.border(c,'bottom','C59B27',8)
 p=h.cell(0,0).paragraphs[0];p.paragraph_format.line_spacing=1.0
 p.add_run().add_picture(str(w.BASE/'assets/rotbeland-logo.png'),width=Mm(13))
 c=h.cell(0,1)
 w.text(c.paragraphs[0],INSTITUTE,8,True)
 w.text(c.add_paragraph(),title,11.5,True)
 if subtitle:w.text(c.add_paragraph(),subtitle,7.5,True)
 w.text(h.cell(0,2).paragraphs[0],week,8)
 w.text(h.cell(0,3).paragraphs[0],date,8)
 return h


def social(w,cell,report=False):
 p=cell.paragraphs[0]
 p.paragraph_format.line_spacing=Pt(12)
 p._p.get_or_add_pPr().append(w.el('bidi',val='0'))
 for name in ['telegram','instagram']:
  p.add_run().add_picture(str(w.BASE/f'assets/{name}.png'),width=Mm(3))
  r=p.add_run(' '+('@' if report else '')+'Rotbelands   ')
  r.font.name=w.FONT;r.font.size=Pt(8)
  rp=r._r.get_or_add_rPr();rp.append(w.el('rtl',val='0'))
  for k in ['ascii','hAnsi','eastAsia','cs']:rp.rFonts.set(qn('w:'+k),w.FONT)


def footer(w,report=False,part=1):
 widths=([70,128] if part==1 else [110,65,23]) if report else [112,90]
 t=w.grid(1,len(widths),'Report Header' if report else 'Planner Header',widths)
 t.rows[0].height=Mm(5)
 for c in t.rows[0].cells:
  for edge in ['bottom','left','right']:w.border(c,edge,val='nil')
  w.border(c,'top','C59B27',4)
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
 if report:
  social(w,t.cell(0,0 if part==1 else 1),True)
  w.text(t.cell(0,1 if part==1 else 0).paragraphs[0],TAGLINE,6.5)
  if part==2:w.text(t.cell(0,2).paragraphs[0],'برگ ۲ از ۲',6.5)
 else:
  w.text(t.cell(0,0).paragraphs[0],'«موفقیت استمرار روزانه در کارهای کوچک است»',7)
  social(w,t.cell(0,1))


def planner_header(w,week,half):
 """Original logo/title/badge, outlined week/date and five separate metadata fields."""
 h=w.grid(1,3,'Planner Header',[15,108,79]);h.rows[0].height=Mm(17)
 for c in h.rows[0].cells:
  w.margins(c,left=.5,right=.5)
  c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for edge in ['top','left','right']:w.border(c,edge,val='nil')
  w.border(c,'bottom','C59B27',6)
 p=h.cell(0,0).paragraphs[0];p.paragraph_format.line_spacing=1.0
 p.add_run().add_picture(str(w.BASE/'assets/rotbeland-logo.png'),width=Mm(13))
 c=h.cell(0,1)
 w.text(c.paragraphs[0],INSTITUTE,8,True)
 title=w.local_table(c,1,[57,49]);title.rows[0].height=Mm(7)
 # local_table tightens the parent spacers, but the institute is real content.
 c.paragraphs[0].paragraph_format.line_spacing=Pt(11)
 for run in c.paragraphs[0].runs:run.font.size=Pt(8)
 w.text(title.cell(0,0).paragraphs[0],'برنامه‌ریزی هفتگی',12,True)
 badge=title.cell(0,1);w.shade(badge,'FCF9EE')
 for edge in ['top','bottom','left','right']:w.border(badge,edge,'B88D1D',4)
 w.text(badge.paragraphs[0],['برگ اول: شنبه تا دوشنبه','برگ دوم: سه‌شنبه تا جمعه'][half],7.5,True,align='center')
 dates=w.local_table(h.cell(0,2),1,[28,2,47]);dates.rows[0].height=Mm(7)
 for col,value in [(0,'هفته: '+w.fa(week+1)+' ـ '+['دوازدهم','دهم'][week]),(2,'تاریخ: '+['۱۸ تا ۲۴ مهر ۱۴۰۵','۲۵ مهر تا ۱ آبان ۱۴۰۵'][week])]:
  cell=dates.cell(0,col);w.shade(cell,'FCF9EE')
  for edge in ['top','bottom','left','right']:w.border(cell,edge,'B88D1D',4)
  w.text(cell.paragraphs[0],value,7.5,align='center')
 w.gap(.7)
 meta=w.grid(1,5,'Planner Header',[42,45,43,38,34]);meta.rows[0].height=Mm(6)
 labels=['نام داوطلب:','رشته / پایه:','مشاور:','هدف ساعت:','هدف تست:']
 values=['','تجربی / '+['دوازدهم','دهم'][week],'','۴۱*','۴۹۵']
 for i,(label,value,width) in enumerate(zip(labels,values,[42,45,43,38,34])):
  cell=meta.cell(0,i);w.margins(cell,left=.5,right=.5)
  for edge in ['top','bottom','left','right']:w.border(cell,edge,'CCCCCC',3,val='single' if edge=='bottom' or (edge=='right' and i<4) else 'nil')
  t=w.local_table(cell,1,[17,width-18]);t.rows[0].height=Mm(4.7)
  w.text(t.cell(0,0).paragraphs[0],label,7)
  w.field(t.cell(0,1),value,7)
