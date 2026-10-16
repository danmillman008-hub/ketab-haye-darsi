"""Editable Word version of the supplied A4 form, with fixed grids and embedded fonts."""
import json, uuid, zipfile, re
from pathlib import Path
from lxml import etree
from ooxml_order import normalize
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

fa=lambda x:str(x).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))
BASE=Path(__file__).resolve().parents[1]
REPO=BASE.parent
plan=json.loads((BASE/'plan.json').read_text())
doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=sec.bottom_margin=Mm(4)
sec.left_margin=sec.right_margin=Mm(4)
sec.footer_distance=Mm(1)
FONT='Vazirmatn'

def el(tag,**attrs):
 e=OxmlElement('w:'+tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 return e

normal=doc.styles['Normal'];normal.font.name=FONT;normal.font.size=Pt(9)
nf=normal.paragraph_format
nf.space_before=nf.space_after=Pt(0);nf.line_spacing=Pt(13)
nf.widow_control=False
# Explicit complex-script defaults avoid theme/locale font substitutions in Word.
rp=normal.element.get_or_add_rPr()
fonts=rp.find(qn('w:rFonts'))
for key in list(fonts.attrib):
 if 'Theme' in key:del fonts.attrib[key]
for key in ['ascii','hAnsi','eastAsia','cs']:fonts.set(qn('w:'+key),FONT)
rp.append(el('szCs',val=18));rp.append(el('lang',val='fa-IR',bidi='fa-IR'))
settings=doc.settings.element
compat=settings.find(qn('w:compat'))
if compat is not None:
 for setting in compat.findall(qn('w:compatSetting')):
  if setting.get(qn('w:name'))=='compatibilityMode':setting.set(qn('w:val'),'15')
for name in ['Planner Header','Planner Day','Planner Exam']:doc.styles.add_style(name,WD_STYLE_TYPE.TABLE)

def text(p,s='',size=9,bold=False,align='start'):
 pp=p._p.get_or_add_pPr()
 for tag in ['bidi','jc']:
  for old in list(pp.findall(qn('w:'+tag))):pp.remove(old)
 pp.append(el('bidi',val='1'));pp.append(el('jc',val=align))
 p.paragraph_format.keep_with_next=False
 p.paragraph_format.widow_control=False
 p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
 p.paragraph_format.line_spacing=Pt(size*1.45)
 for chunk in re.split(r'([☐○])',s):
  if not chunk:continue
  family='DejaVu Sans' if chunk in ['☐','○'] else FONT
  r=p.add_run(chunk);r.bold=bold;r.font.size=Pt(size);r.font.name=family
  rp=r._r.get_or_add_rPr();rp.append(el('rtl',val='1'));rp.append(el('szCs',val=int(size*2)))
  for k in ['ascii','hAnsi','eastAsia','cs']:rp.rFonts.set(qn('w:'+k),family)
  if bold:rp.append(el('bCs',val='1'))
 return p

def border(cell,edge,color='222222',size=4,val='single'):
 pr=cell._tc.get_or_add_tcPr();bs=pr.find(qn('w:tcBorders'))
 if bs is None:bs=el('tcBorders');pr.append(bs)
 for old in list(bs):
  if old.tag==qn('w:'+edge):bs.remove(old)
 bs.append(el(edge,val=val,sz=size,color=color))

def grid(rows,cols,name,widths=None):
 t=doc.add_table(rows=rows,cols=cols);t.style=name;t.autofit=False
 t.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=widths or [202/cols]*cols
 pr=t._tbl.tblPr
 pr.append(el('bidiVisual',val='1'))
 pr.find(qn('w:tblW')).set(qn('w:type'),'dxa');pr.find(qn('w:tblW')).set(qn('w:w'),str(int(Mm(sum(widths)).twips)))
 margins=el('tblCellMar')
 for edge,val in [('top',35),('bottom',25),('left',55),('right',55)]:margins.append(el(edge,w=val,type='dxa'))
 pr.append(margins)
 for col,w in zip(t.columns,widths):col.width=Mm(w)
 for row in t.rows:
  row._tr.get_or_add_trPr().append(el('cantSplit'))
  row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
  for c,w in zip(row.cells,widths):
   c.width=Mm(w)
   for edge in ['top','bottom','left','right']:border(c,edge)
 return t

def shade(cell,color='D1A843'):cell._tc.get_or_add_tcPr().append(el('shd',fill=color))
def gap(mm=1):
 p=doc.add_paragraph();p.paragraph_format.line_spacing=Mm(mm)
 p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
 p.add_run().font.size=Pt(1)

def tight(p):
 p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
 p.paragraph_format.line_spacing=Pt(1)
 p.paragraph_format.keep_with_next=False
 for r in p.runs:r.font.size=Pt(1)


def margins(cell,top=0,bottom=0,left=0,right=0):
 pr=cell._tc.get_or_add_tcPr()
 for old in list(pr.findall(qn('w:tcMar'))):pr.remove(old)
 m=el('tcMar')
 for edge,value in [('top',top),('bottom',bottom),('left',left),('right',right)]:m.append(el(edge,w=round(Mm(value).twips),type='dxa'))
 pr.append(m)


def local_table(cell,rows,widths):
 t=cell.add_table(rows=rows,cols=len(widths));t.autofit=False
 t.alignment=WD_TABLE_ALIGNMENT.CENTER
 pr=t._tbl.tblPr;pr.append(el('bidiVisual',val='1'))
 tw=pr.find(qn('w:tblW'));tw.set(qn('w:type'),'dxa');tw.set(qn('w:w'),str(int(Mm(sum(widths)).twips)))
 mar=el('tblCellMar')
 for side in ['top','bottom','left','right']:mar.append(el(side,w=0,type='dxa'))
 pr.append(mar)
 for col,width in zip(t.columns,widths):col.width=Mm(width)
 for row in t.rows:
  row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
  row._tr.get_or_add_trPr().append(el('cantSplit'))
  for c,width in zip(row.cells,widths):
   c.width=Mm(width);c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
   for edge in ['top','bottom','left','right']:border(c,edge,val='nil')
 for p in cell.paragraphs:tight(p)
 return t


def line(p,edge='bottom',color='9CA3AF',kind='dotted'):
 pr=p._p.get_or_add_pPr()
 for old in list(pr.findall(qn('w:pBdr'))):pr.remove(old)
 bs=el('pBdr');bs.append(el(edge,val=kind,sz=3,space=0,color=color));pr.append(bs)


def field(cell,value='',size=7):
 p=text(cell.paragraphs[0],value,size,align='center');line(p)
 return p


def metric_pair(cell,hours=None,tests=None):
 # Native fixed columns, not runs of padding spaces: hours right, tests left.
 margins(cell,left=.7,right=.7)
 t=local_table(cell,1,[7,12,27.33,7,12]);t.rows[0].height=Mm(4.4)
 text(t.cell(0,0).paragraphs[0],'ساعت:',7)
 field(t.cell(0,1),'' if hours is None else fa(hours))
 text(t.cell(0,3).paragraphs[0],'تست:',7)
 field(t.cell(0,4),'' if tests is None else fa(tests))


def fill(content,metrics,b):
 margins(content,top=.6,bottom=.3,left=1,right=1)
 if b:
  text(content.paragraphs[0],'☐ '+b['subject'],9,True)
  p=text(content.add_paragraph(),b['topic'],9,True);p.paragraph_format.space_before=Pt(4)
  p=text(content.add_paragraph(),b['detail'],8.5);p.paragraph_format.space_before=Pt(4)
 else:text(content.paragraphs[0],'☐ درس و مبحث:',8)
 metric_pair(metrics,b['hours'] if b else None,b['tests'] if b else None)
 border(content,'bottom','CCCCCC',3);border(metrics,'top','CCCCCC',3)


def save(out):
 doc.save(out)
 W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
 R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
 REL='http://schemas.openxmlformats.org/package/2006/relationships'
 CT='http://schemas.openxmlformats.org/package/2006/content-types'
 with zipfile.ZipFile(out) as z:parts={name:z.read(name) for name in z.namelist()}
 font_table=etree.fromstring(parts['word/fontTable.xml'])
 font=next((x for x in font_table if x.get('{'+W+'}name')==FONT),None)
 if font is None:font=etree.SubElement(font_table,'{'+W+'}font',{'{'+W+'}name':FONT})
 rels=etree.Element('{'+REL+'}Relationships',nsmap={None:REL})
 for kind in ['Regular','Bold','Symbols']:
  data=bytearray(((BASE/'assets/word-fonts/DejaVuSymbols.ttf') if kind=='Symbols' else (REPO/'assets/word-fonts'/('Vazirmatn-'+kind+'.ttf'))).read_bytes())
  target=font
  if kind=='Symbols':
   target=next((x for x in font_table if x.get('{'+W+'}name')=='DejaVu Sans'),None)
   if target is None:target=etree.SubElement(font_table,'{'+W+'}font',{'{'+W+'}name':'DejaVu Sans'})
  key=uuid.uuid5(uuid.NAMESPACE_URL,'planner-vazirmatn-'+kind)
  xor=key.bytes[::-1]
  for i in range(32):data[i]^=xor[i%16]
  filename='Vazirmatn-'+kind+'.odttf';rid='rId'+kind
  parts['word/fonts/'+filename]=bytes(data)
  etree.SubElement(target,'{'+W+'}embed'+('Regular' if kind=='Symbols' else kind),{'{'+R+'}id':rid,'{'+W+'}fontKey':'{'+str(key).upper()+'}','{'+W+'}subsetted':'1' if kind=='Symbols' else '0'})
  etree.SubElement(rels,'{'+REL+'}Relationship',Id=rid,Type=R+'/font',Target='fonts/'+filename)
 parts['word/fontTable.xml']=etree.tostring(font_table,xml_declaration=True,encoding='UTF-8',standalone=True)
 parts['word/_rels/fontTable.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
 ct=etree.fromstring(parts['[Content_Types].xml'])
 etree.SubElement(ct,'{'+CT+'}Default',Extension='odttf',ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont')
 parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
 settings=etree.fromstring(parts['word/settings.xml']);settings.append(el('embedTrueTypeFonts'));settings.append(el('saveSubsetFonts',val='0'))
 parts['word/settings.xml']=etree.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
 # Canonical schema order for properties added through low-level XML helpers.
 for filename in ['word/document.xml','word/styles.xml','word/settings.xml']:
  parts[filename]=etree.tostring(normalize(etree.fromstring(parts[filename])),xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
  for name,data in parts.items():z.writestr(name,data)
