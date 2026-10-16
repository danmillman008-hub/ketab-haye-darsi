"""Two blank editable student-report sheets, matching Stitch 13 a4_1/a4_2."""
import importlib
import word_layout as w
from docx.shared import Mm,Pt
from docx.enum.style import WD_STYLE_TYPE

for part in [1,2]:
 if part==2:w=importlib.reload(w)
 doc=w.doc
 sec=doc.sections[0];sec.left_margin=sec.right_margin=Mm(6)
 sec.top_margin=sec.bottom_margin=Mm(6)
 for name in ['Report Header','Report Metadata','Report Day','Report Friday']:
  doc.styles.add_style(name,WD_STYLE_TYPE.TABLE)
 h=w.grid(2,2,'Report Header',[132,66])
 h.rows[0].height=Mm(8);h.rows[1].height=Mm(6)
 for row in h.rows:
  for c in row.cells:
   for edge in ['top','bottom','left','right']:w.border(c,edge,val='nil')
 w.border(h.cell(0,0),'left','D1A843',16)
 w.text(h.cell(0,0).paragraphs[0],'گزارش کار و پایش تست هفتگی داوطلب',13,True)
 w.text(h.cell(0,1).paragraphs[0],'هفته: ..............',9)
 w.text(h.cell(1,0).paragraphs[0],('برگ ۲: سه‌شنبه تا جمعه' if part==2 else ''),8)
 w.text(h.cell(1,1).paragraphs[0],'تاریخ: .... / .... / ۱۴۰....',9)
 for c in h.rows[1].cells:w.border(c,'bottom','B88D1D',10)
 w.gap(1)
 meta=w.grid(1,4,'Report Metadata',[49.5]*4);meta.rows[0].height=Mm(9)
 for c,label in zip(meta.rows[0].cells,['نام داوطلب:','رشته / پایه:','مشاور ناظر:','هدف ساعت / تست:']):
  w.shade(c,'F8F9FA')
  for edge in ['top','bottom','left','right']:w.border(c,edge,'D1D5DB',3)
  w.text(c.paragraphs[0],label,8,True)
  w.text(c.add_paragraph(),'........................................',7)
 w.gap(2)

 def card(cell,title,optional,short):
  if optional:w.shade(cell,'F8F9FA')
  p=w.text(cell.paragraphs[0],title+(' ....................' if optional else ''),8.5,not optional)
  pr=p._p.get_or_add_pPr();bs=w.el('pBdr');bs.append(w.el('bottom',val='single',sz=3,color='D1D5DB'));pr.append(bs)
  for label in ['مبحث:','زمان:']:
   p=w.text(cell.add_paragraph(),label+' ................................',8)
   p.paragraph_format.space_before=Pt(7 if short else 14)
  p=w.text(cell.add_paragraph(),'تست: ..............  (ص: ... غ: ...)',7.5)
  p.paragraph_format.space_before=Pt(7 if short else 14)

 def day(name,friday=False):
  t=w.grid(2 if friday else 3,4,'Report Friday' if friday else 'Report Day',[49.5]*4)
  t.rows[0].height=Mm(6)
  daycell=t.cell(0,0);rating=t.cell(0,1).merge(t.cell(0,3))
  for c in [daycell,rating]:w.shade(c,'B88D1D' if friday else 'D1A843')
  w.text(daycell.paragraphs[0],name,8.5,True)
  w.text(rating.paragraphs[0],'میزان رضایت:  ○ عالی   ○ خوب   ○ متوسط   ○ نیاز به بازنگری',8)
  for ri in range(1,len(t.rows)):
   t.rows[ri].height=Mm(36.5 if part==1 else 30.5)
   for j,c in enumerate(t.rows[ri].cells):
    optional=(ri==2 or friday)
    card(c,'درس اختیاری:' if optional else ['زیست‌شناسی','فیزیک','شیمی','ریاضی'][j],optional,part==2)
  # Gold perimeter, quiet grey inner grid, as in the supplied report form.
  for ri,row in enumerate(t.rows):
   for ci,c in enumerate(row.cells):
    for edge in ['top','bottom','left','right']:w.border(c,edge,'D1D5DB',3)
    if ri==0:w.border(c,'top','B88D1D',8)
    if ri==len(t.rows)-1:w.border(c,'bottom','B88D1D',8)
    if ci==0:w.border(c,'left','B88D1D',8)
    if ci==3:w.border(c,'right','B88D1D',8)
  w.gap(2)
 for name in (['شنبه','یکشنبه','دوشنبه'] if part==1 else ['سه‌شنبه','چهارشنبه','پنج‌شنبه']):day(name)
 if part==2:day('جمعه؛ آزمون، تحلیل و جبرانی',True)
 p=w.text(doc.add_paragraph(),'فرم پایش و سنجش عملکرد تحصیلی هفتگی داوطلبان کنکور سراسری       برگ '+w.fa(part)+' از ۲',7)
 bs=w.el('pBdr');bs.append(w.el('top',val='single',sz=8,color='B88D1D'));p._p.get_or_add_pPr().append(bs)
 doc.core_properties.title='گزارش دانش‌آموز — برگ '+w.fa(part)
 w.save(w.BASE/f'student-report-{part}.docx')
 print('Built report',part)
