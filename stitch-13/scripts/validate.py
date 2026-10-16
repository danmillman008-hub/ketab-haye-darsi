"""Validate new outputs without altering the earlier release."""
import json,zipfile
from pathlib import Path
from bs4 import BeautifulSoup as Soup
from docx import Document
from lxml import etree
import pymupdf
BASE=Path(__file__).resolve().parents[1]
plan=json.loads((BASE/'plan.json').read_text())
assert plan==json.loads((BASE.parent/'plan.json').read_text()), 'Only form changes are intended'
assert len(plan)==12
assert sum(x['tests'] for x in plan)==990 and sum(x['hours'] for x in plan)==82
s=Soup((BASE/'index.html').read_text(),'html.parser')
assert len(s.select('.sheet'))==4
blank=0
for n,sheet in enumerate(s.select('.sheet')):
 template=Soup((BASE/'templates'/('planner-first.html' if n%2==0 else 'planner-second.html')).read_text(),'html.parser')
 for j,sec in enumerate(sheet.select('section')[:3]):
  d=plan[(n//2)*6+(n%2)*3+j]
  cells=sec.select_one('div.grid').find_all('div',recursive=False)
  old=template.select('section')[j].select_one('div.grid').find_all('div',recursive=False)
  for k,b in enumerate(d['boxes']):
   for key in ['subject','topic','detail']:assert b[key] in cells[k].get_text()
  for k in range(len(d['boxes']),6):assert str(cells[k])==str(old[k]);blank+=1
assert blank==18
assert len(pymupdf.open(BASE/'plan-1-aban-1405.pdf'))==4

def word(name):
 with zipfile.ZipFile(BASE/name) as z:
  assert not any(p.startswith('word/media/') for p in z.namelist()),'Word must be editable, not page images'
  assert len([p for p in z.namelist() if p.endswith('.odttf')])==3
  xml=etree.fromstring(z.read('word/document.xml'))
  ns={'w':'http://schemas.openxmlformats.org/wordprocessingml/2006/main'}
  assert all(x.get('{'+ns['w']+'}hRule')=='exact' for x in xml.xpath('//w:trHeight',namespaces=ns))
 return Document(BASE/name)
doc=word('plan-1-aban-1405.docx')
tables=[t for t in doc.tables if t.style.name=='Planner Day'];assert len(tables)==12
for d,t in zip(plan,tables):
 cells=[c for row in [t.rows[1],t.rows[3]] for c in row.cells]
 for i,b in enumerate(d['boxes']):
  for key in ['subject','topic','detail']:assert b[key] in cells[i].text
 for c in cells[len(d['boxes']):]:assert c.text=='☐ درس و مبحث:'
for part in [1,2]:
 doc=word(f'student-report-{part}.docx')
 tables=[t for t in doc.tables if t.style.name in ['Report Day','Report Friday']]
 assert len(tables)==(3 if part==1 else 4)
 count=0;optional=0
 for t in tables:
  assert 'میزان رضایت:' in t.rows[0].cells[1].text
  assert t.rows[0].cells[1].text.count('○')==4
  for row in list(t.rows)[1:]:
   assert len(row.cells)==4
   for c in row.cells:
    assert len(c.tables)==1, 'Each report card must have its own fixed field grid'
    inner=c.tables[0]
    text=' '.join(inner._tbl.xpath('.//w:t/text()'))
    assert all(x in text for x in ['مبحث:','زمان:','تست:','ص:','غ:'])
    assert len(inner.rows)==4
    for ri in [1,2,3]:
     field=inner.cell(ri,1)
     assert field.text=='', 'Writing field must be genuinely blank, not dots or spaces'
     assert field._tc.xpath('.//w:pBdr/w:bottom[@w:val="dotted"]')
    count+=1;optional+=('درس اختیاری:' in text)
 assert count==(24 if part==1 else 28)
 assert optional==(12 if part==1 else 16)
print('PASS: unchanged plan, 4-page PDF, original 18 blank cells, 3 editable RTL Word files, 24+28 report cells and all rating circles.')
