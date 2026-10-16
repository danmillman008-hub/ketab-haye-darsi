"""Editable Word version of the supplied A4 form, with fixed grids and embedded fonts."""
import json, uuid, zipfile, re
from pathlib import Path
from lxml import etree
from docx import Document
from docx.shared import Mm, Pt
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

fa=lambda x:str(x).translate(str.maketrans('0123456789','۰۱۲۳۴۵۶۷۸۹'))
BASE=Path(__file__).resolve().parents[1]
REPO=BASE.parent
plan=json.loads((BASE/'plan.json').read_text())
doc=Document();sec=doc.sections[0]
sec.page_width=Mm(210);sec.page_height=Mm(297)
sec.top_margin=sec.bottom_margin=Mm(4)
sec.left_margin=sec.right_margin=Mm(4)
sec.footer_distance=Mm(1)
FONT='Vazirmatn'

def el(tag,**attrs):
 e=OxmlElement('w:'+tag)
 for k,v in attrs.items():e.set(qn('w:'+k),str(v))
 return e

normal=doc.styles['Normal'];normal.font.name=FONT;normal.font.size=Pt(9)
nf=normal.paragraph_format
nf.space_before=nf.space_after=Pt(0);nf.line_spacing=Pt(13)
nf.widow_control=False
for name in ['Planner Header','Planner Day','Planner Exam']:doc.styles.add_style(name,WD_STYLE_TYPE.TABLE)

def text(p,s='',size=9,bold=False,align='start'):
 pp=p._p.get_or_add_pPr()
 pp.append(el('bidi',val='1'));pp.append(el('jc',val=align))
 p.paragraph_format.keep_with_next=False
 p.paragraph_format.widow_control=False
 p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
 p.paragraph_format.line_spacing=Pt(size*1.45)
 for chunk in re.split(r'([☐○])',s):
  if not chunk:continue
  family='DejaVu Sans' if chunk in ['☐','○'] else FONT
  r=p.add_run(chunk);r.bold=bold;r.font.size=Pt(size);r.font.name=family
  rp=r._r.get_or_add_rPr();rp.append(el('rtl',val='1'));rp.append(el('szCs',val=int(size*2)))
  for k in ['ascii','hAnsi','eastAsia','cs']:rp.rFonts.set(qn('w:'+k),family)
  if bold:rp.append(el('bCs',val='1'))
 return p

def border(cell,edge,color='222222',size=4,val='single'):
 pr=cell._tc.get_or_add_tcPr();bs=pr.find(qn('w:tcBorders'))
 if bs is None:bs=el('tcBorders');pr.append(bs)
 for old in list(bs):
  if old.tag==qn('w:'+edge):bs.remove(old)
 bs.append(el(edge,val=val,sz=size,color=color))

def grid(rows,cols,name,widths=None):
 t=doc.add_table(rows=rows,cols=cols);t.style=name;t.autofit=False
 t.alignment=WD_TABLE_ALIGNMENT.CENTER
 widths=widths or [202/cols]*cols
 pr=t._tbl.tblPr
 pr.append(el('bidiVisual',val='1'))
 pr.find(qn('w:tblW')).set(qn('w:type'),'dxa');pr.find(qn('w:tblW')).set(qn('w:w'),str(int(Mm(sum(widths)).twips)))
 margins=el('tblCellMar')
 for edge,val in [('top',35),('bottom',25),('left',55),('right',55)]:margins.append(el(edge,w=val,type='dxa'))
 pr.append(margins)
 for col,w in zip(t.columns,widths):col.width=Mm(w)
 for row in t.rows:
  row._tr.get_or_add_trPr().append(el('cantSplit'))
  row.height_rule=WD_ROW_HEIGHT_RULE.EXACTLY
  for c,w in zip(row.cells,widths):
   c.width=Mm(w)
   for edge in ['top','bottom','left','right']:border(c,edge)
 return t

def shade(cell,color='D1A843'):cell._tc.get_or_add_tcPr().append(el('shd',fill=color))
def gap(mm=1):
 p=doc.add_paragraph();p.paragraph_format.line_spacing=Mm(mm)
 p.paragraph_format.space_before=p.paragraph_format.space_after=Pt(0)
 p.add_run().font.size=Pt(1)

def fill(content,metrics,b):
 if b:
  text(content.paragraphs[0],'☐ '+b['subject'],9,True)
  p=text(content.add_paragraph(),b['topic'],9,True);p.paragraph_format.space_before=Pt(5)
  p=text(content.add_paragraph(),b['detail'],8.5);p.paragraph_format.space_before=Pt(5)
  text(metrics.paragraphs[0],'ساعت: '+fa(b['hours'])+'                       تست: '+fa(b['tests']),7)
 else:
  text(content.paragraphs[0],'☐ درس و مبحث:',8)
  text(metrics.paragraphs[0],'ساعت: ..............           تست: ..............',7)
 border(content,'bottom','CCCCCC',3);border(metrics,'top','CCCCCC',3)


def save(out):
 doc.save(out)
 W='http://schemas.openxmlformats.org/wordprocessingml/2006/main'
 R='http://schemas.openxmlformats.org/officeDocument/2006/relationships'
 REL='http://schemas.openxmlformats.org/package/2006/relationships'
 CT='http://schemas.openxmlformats.org/package/2006/content-types'
 with zipfile.ZipFile(out) as z:parts={name:z.read(name) for name in z.namelist()}
 font_table=etree.fromstring(parts['word/fontTable.xml'])
 font=next((x for x in font_table if x.get('{'+W+'}name')==FONT),None)
 if font is None:font=etree.SubElement(font_table,'{'+W+'}font',{'{'+W+'}name':FONT})
 rels=etree.Element('{'+REL+'}Relationships',nsmap={None:REL})
 for kind in ['Regular','Bold','Symbols']:
  data=bytearray(((BASE/'assets/word-fonts/DejaVuSymbols.ttf') if kind=='Symbols' else (REPO/'assets/word-fonts'/('Vazirmatn-'+kind+'.ttf'))).read_bytes())
  target=font
  if kind=='Symbols':
   target=next((x for x in font_table if x.get('{'+W+'}name')=='DejaVu Sans'),None)
   if target is None:target=etree.SubElement(font_table,'{'+W+'}font',{'{'+W+'}name':'DejaVu Sans'})
  key=uuid.uuid5(uuid.NAMESPACE_URL,'planner-vazirmatn-'+kind)
  xor=key.bytes[::-1]
  for i in range(32):data[i]^=xor[i%16]
  filename='Vazirmatn-'+kind+'.odttf';rid='rId'+kind
  parts['word/fonts/'+filename]=bytes(data)
  etree.SubElement(target,'{'+W+'}embed'+('Regular' if kind=='Symbols' else kind),{'{'+R+'}id':rid,'{'+W+'}fontKey':'{'+str(key).upper()+'}','{'+W+'}subsetted':'1' if kind=='Symbols' else '0'})
  etree.SubElement(rels,'{'+REL+'}Relationship',Id=rid,Type=R+'/font',Target='fonts/'+filename)
 parts['word/fontTable.xml']=etree.tostring(font_table,xml_declaration=True,encoding='UTF-8',standalone=True)
 parts['word/_rels/fontTable.xml.rels']=etree.tostring(rels,xml_declaration=True,encoding='UTF-8',standalone=True)
 ct=etree.fromstring(parts['[Content_Types].xml'])
 etree.SubElement(ct,'{'+CT+'}Default',Extension='odttf',ContentType='application/vnd.openxmlformats-officedocument.obfuscatedFont')
 parts['[Content_Types].xml']=etree.tostring(ct,xml_declaration=True,encoding='UTF-8',standalone=True)
 settings=etree.fromstring(parts['word/settings.xml']);settings.append(el('embedTrueTypeFonts'));settings.append(el('saveSubsetFonts',val='0'))
 parts['word/settings.xml']=etree.tostring(settings,xml_declaration=True,encoding='UTF-8',standalone=True)
 with zipfile.ZipFile(out,'w',zipfile.ZIP_DEFLATED) as z:
  for name,data in parts.items():z.writestr(name,data)
